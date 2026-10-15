"""LibreOffice で数式を再計算し、エラー(#REF! #DIV/0! など)がないか調べる。

使い方:
    python3 scripts/recalc.py output/売上_編集済.xlsx            # エラー検査
    python3 scripts/recalc.py output/売上_編集済.xlsx --show 20  # 数式の計算結果も20個表示

元のファイルは変更しない(一時フォルダのコピーで計算する)。
openpyxl で保存したファイルは数式の計算結果を持たないが、Excel / Google スプレッドシートで
開くと自動で計算される。LibreOffice で保存し直すと日本語フォントが置き換わるため、上書きはしない。

終了コード: 0 = エラーなし / 1 = 数式エラーあり / 2 = 再計算できなかった
"""

import argparse
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ERROR_VALUES = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}

# 再計算を「常に行う」設定にした使い捨ての LibreOffice プロファイル
PROFILE_XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""

SOFFICE_CANDIDATES = [
    "soffice",
    "libreoffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
]


def find_soffice():
    if os.environ.get("SOFFICE"):
        return os.environ["SOFFICE"]
    for candidate in SOFFICE_CANDIDATES:
        found = shutil.which(candidate) or (candidate if Path(candidate).exists() else None)
        if found:
            return found
    return None


def recalculate(src, workdir, timeout):
    soffice = find_soffice()
    if not soffice:
        print("LibreOffice が見つかりません。インストールするか、環境変数 SOFFICE に soffice の場所を設定してください。")
        sys.exit(2)

    profile = workdir / "profile"
    (profile / "user").mkdir(parents=True)
    (profile / "user" / "registrymodifications.xcu").write_text(PROFILE_XCU, encoding="utf-8")
    outdir = workdir / "out"
    outdir.mkdir()

    cmd = [
        soffice, "--headless", "--norestore",
        f"-env:UserInstallation={profile.resolve().as_uri()}",
        "--convert-to", "xlsx", "--outdir", str(outdir), str(src),
    ]
    try:
        subprocess.run(cmd, capture_output=True, timeout=timeout, check=False)
    except subprocess.TimeoutExpired:
        print(f"再計算が {timeout} 秒以内に終わりませんでした (--timeout で延長できます)")
        sys.exit(2)

    result = outdir / (src.stem + ".xlsx")
    if not result.exists():
        print("再計算に失敗しました (LibreOffice が出力ファイルを作りませんでした)")
        sys.exit(2)
    return result


def main():
    parser = argparse.ArgumentParser(description="数式を再計算してエラーを検出")
    parser.add_argument("file", type=Path)
    parser.add_argument("--timeout", type=int, default=120, help="LibreOffice の待ち時間(秒, 既定 120)")
    parser.add_argument("--show", type=int, default=0, help="数式の計算結果をこの個数だけ表示")
    args = parser.parse_args()

    src = args.file
    if not src.exists():
        sys.exit(f"ファイルが見つかりません: {src}")

    wb = load_workbook(src, keep_vba=src.suffix.lower() == ".xlsm")
    formulas = {}
    external = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas[(ws.title, cell.coordinate)] = cell.value
                    if re.search(r"\[\d+\]", cell.value):
                        external.append(f"{ws.title}!{cell.coordinate}")

    if not formulas:
        print("数式はありません。再計算は不要です。")
        return
    if external:
        print(f"注意: 別ファイルを参照する数式があります({', '.join(external[:5])} など)。")
        print("      参照先ファイルがないと再計算で #REF! / #NAME? になります。値を元ファイルから写すか、ユーザーに確認してください。")

    with tempfile.TemporaryDirectory() as tmp:
        result = recalculate(src.resolve(), Path(tmp), args.timeout)
        values = load_workbook(result, data_only=True)

        errors = {}
        shown = 0
        for (sheet, coord), formula in formulas.items():
            value = values[sheet][coord].value if sheet in values.sheetnames else None
            if isinstance(value, str) and value in ERROR_VALUES:
                errors.setdefault(value, []).append(f"{sheet}!{coord}  {formula}")
            if shown < args.show:
                print(f"{sheet}!{coord}: {formula} → {value}")
                shown += 1

        print(f"数式 {len(formulas)} 個を再計算しました。")
        if errors:
            total = sum(len(v) for v in errors.values())
            print(f"エラー {total} 個:")
            for kind, cells in errors.items():
                print(f"  {kind} ({len(cells)}個)")
                for c in cells[:20]:
                    print(f"    {c}")
                if len(cells) > 20:
                    print(f"    … ほか {len(cells) - 20} 個")
        else:
            print("エラーはありません。")

    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()
