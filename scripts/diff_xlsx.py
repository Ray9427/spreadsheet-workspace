"""2つの Excel / CSV ファイルを比べて、変わったセルを一覧表示する。

使い方:
    python3 scripts/diff_xlsx.py input/売上.xlsx output/売上_編集済.xlsx
    python3 scripts/diff_xlsx.py 旧.csv 新.csv --max 500

数式は計算結果ではなく式そのもの(=SUM(...) など)を比べる。
"""

import argparse
import csv
import datetime as dt
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def fmt(value):
    if value is None:
        return "(空)"
    if isinstance(value, dt.datetime) and not (value.hour or value.minute or value.second):
        return value.date().isoformat()
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()
    text = str(value).replace("\n", "\\n")
    return text if len(text) <= 50 else text[:47] + "..."


def load_cells(path):
    """{シート名: {(行, 列): 値}} を返す。CSV は1シート扱い。"""
    suffix = path.suffix.lower()
    if suffix in (".csv", ".tsv", ".txt"):
        delimiter = "\t" if suffix == ".tsv" else ","
        for encoding in ("utf-8-sig", "cp932"):
            try:
                with open(path, newline="", encoding=encoding) as f:
                    rows = list(csv.reader(f, delimiter=delimiter))
                break
            except UnicodeDecodeError:
                continue
        else:
            sys.exit(f"文字コードを判定できませんでした: {path}")
        cells = {(r, c): v for r, row in enumerate(rows, 1) for c, v in enumerate(row, 1) if v != ""}
        return {"(CSV)": cells}

    wb = load_workbook(path, keep_vba=suffix == ".xlsm")
    sheets = {}
    for ws in wb.worksheets:
        sheets[ws.title] = {
            (cell.row, cell.column): cell.value
            for row in ws.iter_rows()
            for cell in row
            if cell.value is not None
        }
    return sheets


def main():
    parser = argparse.ArgumentParser(description="2つの表ファイルの差分を表示")
    parser.add_argument("old", type=Path)
    parser.add_argument("new", type=Path)
    parser.add_argument("--max", type=int, default=200, help="表示する差分の最大件数 (既定 200)")
    args = parser.parse_args()

    for p in (args.old, args.new):
        if not p.exists():
            sys.exit(f"ファイルが見つかりません: {p}")

    old, new = load_cells(args.old), load_cells(args.new)
    print(f"比較: {args.old}  →  {args.new}")

    added_sheets = [s for s in new if s not in old]
    removed_sheets = [s for s in old if s not in new]
    if added_sheets:
        print(f"追加されたシート: {', '.join(added_sheets)}")
    if removed_sheets:
        print(f"削除されたシート: {', '.join(removed_sheets)}")
    if list(new) != list(old) and not added_sheets and not removed_sheets:
        print(f"シートの順番が変わりました: {', '.join(old)} → {', '.join(new)}")

    total = 0
    printed = 0
    for sheet in [s for s in old if s in new]:
        a, b = old[sheet], new[sheet]
        changed = sorted(k for k in a.keys() | b.keys() if a.get(k) != b.get(k))
        if not changed:
            continue
        total += len(changed)
        print()
        print(f"=== {sheet}: {len(changed)} セル変更 ===")
        for r, c in changed:
            if printed >= args.max:
                break
            print(f"  {get_column_letter(c)}{r}: {fmt(a.get((r, c)))} → {fmt(b.get((r, c)))}")
            printed += 1

    print()
    if total == 0 and not added_sheets and not removed_sheets:
        print("セルの値・数式に違いはありません。")
    else:
        print(f"合計 {total} セル変更" + (f"(うち {printed} 件を表示。--max で変更)" if printed < total else ""))


if __name__ == "__main__":
    main()
