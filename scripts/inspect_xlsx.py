"""Excel / CSV ファイルの中身を、セル番地つきで一覧表示する。

使い方:
    python3 scripts/inspect_xlsx.py input/売上.xlsx
    python3 scripts/inspect_xlsx.py input/売上.xlsx --sheet 4月 --rows 50
    python3 scripts/inspect_xlsx.py input/data.csv

数式セルは「=式 → 計算結果」の形で表示する(計算結果はファイルに保存されている値)。
"""

import argparse
import csv
import datetime as dt
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def fmt(value):
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%d %H:%M") if (value.hour or value.minute) else value.strftime("%Y-%m-%d")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float):
        return f"{value:.10g}"
    text = str(value).replace("\n", "\\n")
    return text if len(text) <= 60 else text[:57] + "..."


def read_csv_rows(path):
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    for encoding in ("utf-8-sig", "cp932"):
        try:
            with open(path, newline="", encoding=encoding) as f:
                return list(csv.reader(f, delimiter=delimiter)), encoding
        except UnicodeDecodeError:
            continue
    sys.exit(f"文字コードを判定できませんでした: {path}")


def inspect_csv(path, max_rows, max_cols):
    rows, encoding = read_csv_rows(path)
    width = max((len(r) for r in rows), default=0)
    print(f"=== CSV: {path.name} ({len(rows)}行 × {width}列, 文字コード {encoding}) ===")
    for r, row in enumerate(rows[:max_rows], start=1):
        cells = [f"{get_column_letter(c)}: {fmt(v)}" for c, v in enumerate(row[:max_cols], start=1) if v != ""]
        if cells:
            print(f"{r:>4} | " + " | ".join(cells))
    if len(rows) > max_rows:
        print(f"  … 残り {len(rows) - max_rows} 行は省略 (--rows で変更)")


def inspect_xlsx(path, sheet_name, max_rows, max_cols):
    keep_vba = path.suffix.lower() == ".xlsm"
    wb = load_workbook(path, keep_vba=keep_vba)  # 数式
    wb_values = load_workbook(path, data_only=True, keep_vba=keep_vba)  # 保存されている計算結果

    print(f"ファイル: {path.name}")
    print(f"シート一覧: {', '.join(wb.sheetnames)}")
    if wb.defined_names:
        print(f"名前付き範囲: {', '.join(wb.defined_names.keys())}")

    targets = [sheet_name] if sheet_name else wb.sheetnames
    for name in targets:
        if name not in wb.sheetnames:
            sys.exit(f"シート「{name}」がありません。候補: {', '.join(wb.sheetnames)}")
        ws, ws_v = wb[name], wb_values[name]

        formula_count = sum(
            1 for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")
        )
        print()
        print(f"=== シート: {name} (使用範囲 {ws.dimensions}, {ws.max_row}行 × {ws.max_column}列) ===")
        info = [f"数式 {formula_count}個"]
        if ws.sheet_state != "visible":
            info.append(f"状態 {ws.sheet_state}")
        if ws.freeze_panes:
            info.append(f"ウィンドウ枠固定 {ws.freeze_panes}")
        if ws.auto_filter.ref:
            info.append(f"フィルター {ws.auto_filter.ref}")
        if ws.merged_cells.ranges:
            merged = [str(r) for r in ws.merged_cells.ranges]
            info.append(f"結合セル {', '.join(merged[:10])}" + (" …" if len(merged) > 10 else ""))
        if ws.data_validations.dataValidation:
            info.append(f"入力規則 {len(ws.data_validations.dataValidation)}件")
        if len(ws.conditional_formatting):
            info.append(f"条件付き書式 {len(ws.conditional_formatting)}件")
        if ws.tables:
            info.append(f"テーブル {', '.join(f'{t}({ws.tables[t].ref})' for t in ws.tables)}")
        print(" / ".join(info))

        uncomputed = 0
        for r in range(1, min(ws.max_row, max_rows) + 1):
            cells = []
            for c in range(1, min(ws.max_column, max_cols) + 1):
                value = ws.cell(r, c).value
                if value is None:
                    continue
                col = get_column_letter(c)
                if isinstance(value, str) and value.startswith("="):
                    cached = ws_v.cell(r, c).value
                    if cached is None:
                        uncomputed += 1
                    cells.append(f"{col}: {value} → {fmt(cached) if cached is not None else '(未計算)'}")
                else:
                    cells.append(f"{col}: {fmt(value)}")
            if cells:
                print(f"{r:>4} | " + " | ".join(cells))
        if ws.max_row > max_rows:
            print(f"  … {max_rows + 1}行目以降は省略 (--rows で変更)")
        if ws.max_column > max_cols:
            print(f"  … {get_column_letter(max_cols + 1)}列以降は省略 (--cols で変更)")
        if uncomputed:
            print(f"  ※ 未計算の数式が {uncomputed} 個あります。値を確かめるには scripts/recalc.py を使ってください。")


def main():
    parser = argparse.ArgumentParser(description="Excel / CSV の中身をセル番地つきで表示")
    parser.add_argument("file", type=Path)
    parser.add_argument("--sheet", help="このシートだけ表示")
    parser.add_argument("--rows", type=int, default=30, help="表示する最大行数 (既定 30)")
    parser.add_argument("--cols", type=int, default=20, help="表示する最大列数 (既定 20)")
    args = parser.parse_args()

    if not args.file.exists():
        sys.exit(f"ファイルが見つかりません: {args.file}")
    if args.file.suffix.lower() in (".csv", ".tsv", ".txt"):
        inspect_csv(args.file, args.rows, args.cols)
    else:
        inspect_xlsx(args.file, args.sheet, args.rows, args.cols)


if __name__ == "__main__":
    main()
